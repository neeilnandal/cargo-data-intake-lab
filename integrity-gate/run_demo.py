"""Run IntegrityGate on two weekly snapshots of each sample dataset.

Writes:
  data/      the synthetic inputs (CSV), so every result can be traced
  outputs/   exception_register.xlsx, pareto.png, summary.md

Usage:  python run_demo.py
"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from integritygate import samples
from integritygate.engine import carry_forward, run_pack, summarise
from integritygate.rules import (
    CLIENT_DATA,
    PROJECT_FINANCE,
    RECONCILIATION,
    build_reconciliation_frame,
    rule_catalogue,
)

ROOT = Path(__file__).parent
DATA, OUT = ROOT / "data", ROOT / "outputs"

PACK_LABELS = {
    "client_data": "Client data (CSO Data Integrity)",
    "reconciliation": "Core vs report reconciliation (Operations)",
    "project_finance": "Project finance facilities (GCIB)",
}


def run_all() -> tuple[pd.DataFrame, pd.DataFrame, dict[str, int]]:
    DATA.mkdir(exist_ok=True)
    OUT.mkdir(exist_ok=True)

    c1, c2 = samples.client_data()
    (core1, rep1), (core2, rep2) = samples.reconciliation()
    p1, p2 = samples.project_finance()

    for name, df in {
        "client_master_wk1": c1, "client_master_wk2": c2,
        "core_accounts_wk1": core1, "report_accounts_wk1": rep1,
        "core_accounts_wk2": core2, "report_accounts_wk2": rep2,
        "pf_facilities_wk1": p1, "pf_facilities_wk2": p2,
    }.items():
        df.to_csv(DATA / f"{name}.csv", index=False)

    r1 = build_reconciliation_frame(core1, rep1)
    r2 = build_reconciliation_frame(core2, rep2)

    week1 = pd.concat([
        run_pack(CLIENT_DATA, c1, samples.WEEK1),
        run_pack(RECONCILIATION, r1, samples.WEEK1),
        run_pack(PROJECT_FINANCE, p1, samples.WEEK1),
    ], ignore_index=True)
    week2_raw = pd.concat([
        run_pack(CLIENT_DATA, c2, samples.WEEK2),
        run_pack(RECONCILIATION, r2, samples.WEEK2),
        run_pack(PROJECT_FINANCE, p2, samples.WEEK2),
    ], ignore_index=True)
    week2 = carry_forward(week1, week2_raw)

    checked = {"client_data": len(c2), "reconciliation": len(r2), "project_finance": len(p2)}
    return week1, week2, checked


SHORT_NAMES = {
    "CD01": "No LEI", "CD02": "LEI check digits", "CD03": "Duplicate client", "CD04": "KYC review overdue",
    "CD05": "Closed CIF, live account", "CD06": "No signatory", "CD07": "Bad country code",
    "RC01": "Missing in report", "RC02": "Missing in core", "RC03": "Status mismatch",
    "RC04": "Balance mismatch", "RC05": "Currency mismatch",
    "PF01": "Overdrawn", "PF02": "DSCR below covenant", "PF03": "Covenant test overdue",
    "PF04": "Insurance expired", "PF05": "No E&S category", "PF06": "Maturity before close",
    "PF07": "SPV LEI missing",
}


def pareto_chart(register: pd.DataFrame, path: Path) -> None:
    from matplotlib.ticker import MaxNLocator

    live = register[register["status"] != "RESOLVED"]
    surface, ink, muted, grid, bar = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df", "#2a78d6"
    fig, axes = plt.subplots(1, 3, figsize=(14, 4.6), facecolor=surface)
    for ax, pack in zip(axes, PACK_LABELS):
        counts = live[live["pack"] == pack].groupby(["rule_code", "reason"]).size().sort_values()
        labels = [f"{SHORT_NAMES.get(code, code)}  {code}" for code, _ in counts.index]
        ax.xaxis.set_major_locator(MaxNLocator(integer=True))
        ax.set_facecolor(surface)
        ax.barh(labels, counts.values, color=bar, height=0.6)
        for y, v in enumerate(counts.values):
            ax.text(v + 0.15, y, str(v), va="center", fontsize=9, color=ink)
        ax.set_title(PACK_LABELS[pack], fontsize=10.5, color=ink, loc="left", pad=10)
        ax.tick_params(colors=muted, labelsize=9, length=0)
        ax.xaxis.grid(True, color=grid, linewidth=0.8)
        ax.set_axisbelow(True)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.spines["bottom"].set_color(grid)
        ax.set_xlim(0, max(counts.values) * 1.2 if len(counts) else 1)
        top3 = round(100 * counts.values[-3:].sum() / counts.sum()) if len(counts) else 0
        ax.set_xlabel(f"Open breaks ({counts.sum()} in total; top 3 causes = {top3}%)",
                      fontsize=8.5, color=muted, labelpad=8)
    fig.suptitle("Open breaks by reason code, week of 26 Oct 2026 (synthetic data)",
                 fontsize=12, color=ink, x=0.01, ha="left")
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    fig.savefig(path, dpi=160, facecolor=surface)
    plt.close(fig)


def write_workbook(summary: pd.DataFrame, register: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        register.to_excel(xw, sheet_name="Exception register", index=False)
        rule_catalogue().to_excel(xw, sheet_name="Rule catalogue", index=False)

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F3864")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.properties.creator = "Neeil Nandal"
    wb.properties.lastModifiedBy = "Neeil Nandal"
    wb.properties.title = "IntegrityGate exception register (synthetic data)"
    wb.save(path)


def write_summary_md(summary: pd.DataFrame, path: Path) -> None:
    lines = ["# IntegrityGate run summary (synthetic data, week of 26 Oct 2026)", ""]
    for _, s in summary.iterrows():
        lines += [
            f"## {PACK_LABELS[s['pack']]}",
            f"- Records checked: {s['records_checked']}; clean: {s['records_clean']} ({s['clean_rate_pct']}%)",
            f"- Open breaks: {s['open_breaks']} (new this week: {s['new_breaks']}; high severity: {s['high_severity_open']})",
            f"- Resolved since last run: {s['resolved_since_last_run']}",
            f"- Median age of open breaks: {s['median_age_days_open']:.0f} days",
            f"- Breaks with a reason code: {s['reason_coded_pct']}%",
            f"- Top cause: {s['top_cause']}",
            "",
        ]
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    week1, week2, checked = run_all()
    summary = summarise(week2, checked)
    write_workbook(summary, week2, OUT / "exception_register.xlsx")
    pareto_chart(week2, OUT / "pareto.png")
    write_summary_md(summary, OUT / "summary.md")
    print(summary.to_string(index=False))


if __name__ == "__main__":
    main()
